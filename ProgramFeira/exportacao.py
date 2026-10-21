"""Exporta uma cópia da última visualização aplicada, sem consultar ou alterar o banco."""
import csv
import os
import tempfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from ProgramFeira.database.banco import converter_data_salva
from ProgramFeira.resumos import calcular_indicadores

CABECALHO = ('ID', 'Data', 'Dia da semana', 'Dinheiro', 'Pix', 'Cartão', 'Total')
DIAS = ('Segunda-feira', 'Terça-feira', 'Quarta-feira', 'Quinta-feira', 'Sexta-feira', 'Sábado', 'Domingo')


def numero(valor):
    return Decimal(str(valor)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def texto_seguro(valor):
    texto = str(valor)
    return "'" + texto if texto.lstrip().startswith(('=', '+', '-', '@')) else texto


def montar_relatorio(vendas, filtro, comparacao):
    vendas = tuple(tuple(v) for v in vendas)
    indicadores = calcular_indicadores(vendas)
    totais = [sum((numero(v[i]) for v in vendas), Decimal(0)) for i in (2, 3, 4)]
    resumo = [('Filtro aplicado', filtro), ('Registros / dias de feira', len(vendas)),
              ('Dinheiro', totais[0]), ('Pix', totais[1]), ('Cartão', totais[2]),
              ('Total geral', indicadores['total']), ('Média por dia de feira', numero(indicadores['media'])),
              ('Contagem', 'Inclui os registros com pagamentos zerados')]
    for nome, percentual in zip(('Dinheiro', 'Pix', 'Cartão'), indicadores['participacoes']):
        resumo.append((f'Participação {nome}', f'{percentual:.1f}%'.replace('.', ',') if indicadores['total'] else 'Indisponível: total zero'))
    resumo.append(('Comparação com período anterior', comparacao))
    return {'vendas': vendas, 'resumo': tuple(resumo)}


def linhas_vendas(relatorio):
    for id_venda, data, dinheiro, pix, cartao in relatorio['vendas']:
        dia = converter_data_salva(data)
        valores = tuple(numero(v) for v in (dinheiro, pix, cartao))
        yield (id_venda, dia or str(data), DIAS[dia.weekday()] if dia else 'Data inválida', *valores, sum(valores))


def celula_csv(valor):
    if isinstance(valor, Decimal):
        return f'{valor:.2f}'.replace('.', ',')
    if hasattr(valor, 'strftime'):
        return valor.strftime('%d/%m/%Y')
    return texto_seguro(valor)


def escrever_csv(caminho, relatorio):
    with open(caminho, 'w', encoding='utf-8-sig', newline='') as arquivo:
        writer = csv.writer(arquivo, delimiter=';')
        writer.writerow(CABECALHO)
        writer.writerows(tuple(celula_csv(v) for v in linha) for linha in linhas_vendas(relatorio))
        writer.writerow([])
        writer.writerow(['Resumo', 'Valor'])
        writer.writerows((texto_seguro(k), celula_csv(v)) for k, v in relatorio['resumo'])


def escrever_excel(caminho, relatorio):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as erro:
        raise RuntimeError('Para exportar Excel, instale as dependências: py -m pip install -r requirements.txt. A exportação CSV continua disponível.') from erro
    wb = Workbook()
    vendas = wb.active
    vendas.title = 'Vendas'
    vendas.append(CABECALHO)
    for linha in linhas_vendas(relatorio):
        # Datas e valores são células reais, prontas para cálculos no Excel.
        vendas.append([str(v) if isinstance(v, Decimal) else v for v in linha])
        for coluna in range(4, 8):
            celula = vendas.cell(vendas.max_row, coluna)
            celula.value = float(linha[coluna - 1])
            celula.number_format = '#,##0.00'
        celula = vendas.cell(vendas.max_row, 2)
        if not isinstance(linha[1], str):
            celula.number_format = 'dd/mm/yyyy'
        else:
            celula.data_type = 's'
    vendas.freeze_panes = 'A2'
    vendas.auto_filter.ref = vendas.dimensions
    resumo = wb.create_sheet('Resumo')
    resumo.append(('Indicador', 'Valor'))
    for nome, valor in relatorio['resumo']:
        resumo.append((nome, float(valor) if isinstance(valor, Decimal) else valor))
        celula = resumo.cell(resumo.max_row, 2)
        if isinstance(valor, str):
            celula.data_type = 's'
        elif isinstance(valor, Decimal):
            celula.number_format = '#,##0.00'
    for sheet in (vendas, resumo):
        for celula in sheet[1]:
            celula.font = Font(bold=True)
    for col, width in zip('ABCDEFG', (10, 15, 22, 18, 18, 18, 18)):
        vendas.column_dimensions[col].width = width
    resumo.column_dimensions['A'].width = 38
    resumo.column_dimensions['B'].width = 100
    try:
        wb.save(caminho)
    finally:
        wb.close()


def exportar(caminho, relatorio, formato):
    caminho = Path(caminho)
    escritores = {'csv': escrever_csv, 'xlsx': escrever_excel}
    if formato not in escritores or caminho.suffix.lower() != '.' + formato:
        raise ValueError('Escolha um arquivo com a extensão correspondente ao formato exportado.')
    fd, temporario = tempfile.mkstemp(prefix='.programfeira-', suffix='.' + formato, dir=caminho.parent)
    os.close(fd)
    try:
        escritores[formato](temporario, relatorio)
        os.replace(temporario, caminho)
    finally:
        Path(temporario).unlink(missing_ok=True)
