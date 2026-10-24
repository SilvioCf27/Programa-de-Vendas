"""Entrada do pacote portátil e verificação isolada do executável."""
import json
import os
from pathlib import Path
import sys
import tempfile
import traceback


def verificar_pacote(relatorio):
    # O teste do binário nunca usa o banco pessoal, nem o caminho do ambiente.
    from ProgramFeira.versao import VERSAO
    resultado = {'versao': VERSAO, 'sucesso': False, 'plataforma': sys.platform, 'empacotado': bool(getattr(sys, 'frozen', False))}
    try:
        with tempfile.TemporaryDirectory(prefix='gestao-feira-pacote-') as pasta:
            os.environ['PROGRAMFEIRA_DATA_DIR'] = pasta
            import tkinter as tk
            original = tk.Tk.mainloop
            tk.Tk.mainloop = lambda self, *args, **kwargs: None
            try:
                from ProgramFeira import main as app
            finally:
                tk.Tk.mainloop = original
            try:
                app.janela.update()
                for tema in ('claro', 'escuro', 'grafite', 'areia', 'aurora'):
                    app.alternar_tema(tema)
                    app.janela.update()
                    assert app.tema_atual.get() == tema
                configuracoes = app.abrir_configuracoes()
                app.janela.update()
                assert len(configuracoes.acoes) == 7
                assert 'Verificar integridade do backup…' in configuracoes.acoes
                configuracoes.destroy()
                from ProgramFeira.interface.fontes import PASTA_FONTES
                from ProgramFeira.database import banco
                from ProgramFeira import financeiro, backup_drive
                from ProgramFeira.exportacao import exportar, montar_relatorio
                from openpyxl import load_workbook
                vendas = banco.listar_vendas()
                assert len(vendas) == 0, 'O pacote público contém vendas pessoais'
                assert banco.CAMINHO_BANCO.parent == Path(pasta).resolve(), 'Banco fora da pasta temporária'
                fontes = list(PASTA_FONTES.rglob('*.ttf'))
                assert len(fontes) == 4, 'Fontes ausentes no pacote'
                banco.cadastrar_venda('2035-01-01', 10, 20, 30)
                assert banco.buscar_venda_por_data('2035-01-01')[2:] == (10, 20, 30)
                financeiro.salvar_lancamento('despesa', '2035-01-01', 'Teste do pacote', '12,34')
                assert financeiro.primeira_fatura('2035-01-09') == '2035-02-09'
                financeiro.salvar_lancamento('despesa', '2035-10-09', 'Compra no fechamento', '13', 'Cartão de crédito', data_compra='2035-10-09')
                automatica = next(r for r in financeiro.listar_lancamentos('2035-01-01','2035-12-31') if r['descricao'] == 'Compra no fechamento')
                assert (automatica['data_compra'], automatica['data']) == ('2035-10-09','2035-11-09')
                financeiro.salvar_lancamento('despesa', '2035-05-09', 'Compra parcelada fictícia', '100,00', 'Cartão de crédito', parcelas=5)
                parcelas = sorted((r for r in financeiro.listar_lancamentos('2035-05-01', '2035-09-30') if r['grupo']), key=lambda r: r['parcela'])
                assert [r['valor_centavos'] for r in parcelas] == [2000] * 5
                compra = financeiro.compra_do_lancamento(parcelas[3]['id'])
                assert compra['valor_centavos'] == 10000 and compra['data'] == '2035-05-09'
                financeiro.salvar_lancamento('despesa', compra['data'], 'Compra corrigida fictícia', '150,01', 'Cartão de crédito', parcelas=5, identificador=parcelas[3]['id'], esperados=compra['registros'])
                corrigida = financeiro.compra_do_lancamento(parcelas[3]['id'])
                assert [r['valor_centavos'] for r in corrigida['registros']] == [3000, 3000, 3000, 3000, 3001]
                popup_compra = app.painel_financeiro.formulario('despesa', corrigida['registros'][3])
                app.janela.update()
                assert popup_compra.winfo_children()[0].grid_slaves(row=2, column=1)[0].get() == '150,01'
                popup_compra.destroy()
                popup_taxas = app.painel_financeiro.formulario_taxas()
                app.janela.update()
                assert len(popup_taxas.campos_taxas) == 7
                popup_taxas.destroy()
                from ProgramFeira.interface.configuracao_fechamento import ConfiguracaoFechamento
                popup_fechamento = ConfiguracaoFechamento(app.janela, {'versao': 0, 'historico': []}, lambda *args: None)
                app.janela.update()
                assert popup_fechamento.previa.get()
                popup_fechamento.destroy()

                destino = Path(pasta) / 'Drive-teste'
                destino.mkdir()
                copia = backup_drive.executar(banco.CAMINHO_BANCO, backup_drive.configurar(banco.CAMINHO_BANCO, destino))
                assert copia['arquivo'].exists()
                rel = montar_relatorio(vendas, 'Verificação do pacote', '')
                exportar(Path(pasta) / 'teste.csv', rel, 'csv')
                exportar(Path(pasta) / 'teste.xlsx', rel, 'xlsx')
                livro = load_workbook(Path(pasta) / 'teste.xlsx')
                assert livro['Vendas'].max_row == 1
                livro.close()
                # Verifica os módulos de nuvem e a DPAPI dentro do binário, sem rede.
                from ProgramFeira import nuvem, sessao_windows, sincronizacao
                config = nuvem.validar_config('https://teste-pacote.supabase.co', 'sb_publishable_somente_teste')
                sessao = nuvem.Sessao('conta-ficticia', 'teste@example.invalid', 'acesso-ficticio', 'renovavel-ficticio', 1)
                sessao_protegida = False
                if sys.platform == 'win32':
                    assert sessao_windows.DISPONIVEL
                    sessao_windows.salvar(pasta, config['url'], sessao)
                    restaurada = sessao_windows.carregar(pasta, config['url'])
                    assert restaurada.refresh_token == sessao.refresh_token
                    assert restaurada.token == ''
                    sessao_windows.remover(pasta)
                    assert sessao_windows.carregar(pasta, config['url']) is None
                    sessao_protegida = True
                assert not sincronizacao.vinculado(banco.CAMINHO_BANCO)
                assert app.conexao_nuvem.agendador.INTERVALO == 60000
                from ProgramFeira.recuperacao import testar_recuperacao
                assert testar_recuperacao(copia['arquivo'])['sucesso']
                import sqlite3
                from contextlib import closing
                from ProgramFeira.database.armazenamento import restaurar_backup
                adulterado = Path(pasta) / 'backup-adulterado.db'
                adulterado.write_bytes(banco.CAMINHO_BANCO.read_bytes())
                with closing(sqlite3.connect(adulterado)) as conexao, conexao:
                    conexao.execute('CREATE TRIGGER adulterado AFTER INSERT ON vendas BEGIN DELETE FROM vendas; END')
                antes_restauracao = banco.CAMINHO_BANCO.read_bytes()
                recusado = False
                try:
                    restaurar_backup(banco.CAMINHO_BANCO, adulterado)
                except ValueError:
                    recusado = True
                assert recusado and banco.CAMINHO_BANCO.read_bytes() == antes_restauracao
                resultado['sqlite'] = sqlite3.sqlite_version
                from ProgramFeira.tendencias import evolucao
                assert evolucao(banco.listar_vendas())[0]['quantidade'] == 1
                from ProgramFeira.interface.conflitos import descrever
                assert 'excluído' in descrever(None)
                resultado.update(sucesso=True, registros_iniciais=len(vendas), fontes=len(fontes),
                                 cadastro=True, despesas=True, backup=True, csv=True, excel=True, configuracoes=True, temas=True, taxas_semanais=True, fechamento_configuravel=True,
                                 compra_total=True, edicao_compra=True, fatura_automatica=True, backup_seguranca=True, banco_isolado=True, modulos_nuvem=True, sessao_protegida=sessao_protegida, recuperacao=True, evolucao=True, conflitos=True)
            finally:
                app.janela.destroy()
    except Exception:
        resultado['erro'] = traceback.format_exc()
    finally:
        Path(relatorio).write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if resultado['sucesso'] else 1


def iniciar():
    try:
        from ProgramFeira import main  # PyInstaller descobre os módulos estaticamente.
        return 0
    except Exception:
        erro = traceback.format_exc()
        from ProgramFeira.database.armazenamento import pasta_dados
        log = pasta_dados() / 'logs/inicializacao.log'
        try:
            log.parent.mkdir(parents=True, exist_ok=True)
            with log.open('a', encoding='utf-8') as arquivo:
                arquivo.write(erro + '\n')
            mensagem = f'Não foi possível abrir o Gestão da Feira.\n\nDetalhes salvos em:\n{log}\n\n' + erro.splitlines()[-1]
        except OSError:
            mensagem = 'Não foi possível abrir o Gestão da Feira.\n\n' + erro.splitlines()[-1]
        if sys.platform == 'win32':
            import ctypes
            ctypes.windll.user32.MessageBoxW(None, mensagem, 'Gestão da Feira - erro na abertura', 0x10)
        else:
            print(mensagem, file=sys.stderr)
        return 1


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--verificar-pacote':
        sys.exit(verificar_pacote(sys.argv[2]))
    sys.exit(iniciar())
