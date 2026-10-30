"""Smoke do pacote Qt em diretório temporário, sem rede nem dados pessoais."""
from contextlib import closing
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import traceback
from unittest.mock import patch

from reforma_desktop.versao import VERSAO


def verificar(relatorio):
    resultado = dict(versao=VERSAO, sucesso=False, plataforma=sys.platform,
        empacotado=bool(getattr(sys, 'frozen', False)), interface='PySide6')
    janela = None
    try:
        with tempfile.TemporaryDirectory(prefix='feira-pacote-qt-') as temp, \
                patch.dict(os.environ, {'PROGRAMFEIRA_DATA_DIR': temp}), \
                patch('urllib.request.OpenerDirector.open', side_effect=AssertionError('Rede proibida no smoke')):
            from reforma_desktop.producao import conferir_runtime, cliente_configurado
            conferir_runtime()
            from reforma_desktop.aplicacao import preparar_interface, Janela
            from reforma_desktop.vendas import ServicoVendas
            from ProgramFeira.database import banco
            from ProgramFeira.database.armazenamento import restaurar_backup
            from ProgramFeira import financeiro, sessao_windows, nuvem
            from ProgramFeira.exportacao import exportar
            from ProgramFeira.recuperacao import testar_recuperacao, _conteudo
            from openpyxl import load_workbook
            pasta = Path(temp)
            app = preparar_interface()
            servico = ServicoVendas()
            financeiro.inicializar()
            assert Path(banco.CAMINHO_BANCO).resolve().parent == pasta.resolve()
            assert banco.listar_vendas() == []
            janela = Janela(servico, pasta, cliente_nuvem=cliente_configurado(pasta), producao=True)
            janela.atualizador.parar()
            janela.show(); app.processEvents()
            assert not janela.nuvem.cliente.url
            assert janela.atualizador.real
            servico.salvar('2035-01-01', '10', '20', '30')
            financeiro.salvar_lancamento('despesa', '2035-05-09', 'Fictício', '100,01', 'Cartão de crédito', parcelas=3)
            registros = financeiro.listar_lancamentos('2035-01-01', '2035-12-31')
            assert sum(r['valor_centavos'] for r in registros) == 10001
            for area in ('Vendas', 'Análises', 'Finanças'):
                janela.mostrar(area); app.processEvents()
            config = janela.configuracoes(); app.processEvents(); config.close()
            painel = janela.atualizacoes(); app.processEvents()
            assert not painel.cenario.isVisible()
            painel.close()
            janela.limpar()
            exportar(pasta/'teste.xlsx', janela.resultado['relatorio'], 'xlsx')
            livro = load_workbook(pasta/'teste.xlsx')
            assert livro['Vendas'].max_row == 2
            livro.close()
            copia = banco.criar_backup()
            assert testar_recuperacao(copia)['sucesso']
            adulterado = pasta/'adulterado.db'
            adulterado.write_bytes(Path(banco.CAMINHO_BANCO).read_bytes())
            with closing(sqlite3.connect(adulterado)) as c, c:
                c.execute('CREATE TRIGGER adulterado AFTER INSERT ON vendas BEGIN DELETE FROM vendas; END')
            antes = _conteudo(banco.CAMINHO_BANCO)
            try:
                restaurar_backup(banco.CAMINHO_BANCO, adulterado)
            except ValueError:
                pass
            else:
                raise AssertionError('Backup adulterado aceito')
            assert _conteudo(banco.CAMINHO_BANCO) == antes
            raiz = Path(__file__).resolve().parents[1]
            assert (raiz/'prototipos/reforma/recursos/logo.svg').is_file()
            assert list((raiz/'ProgramFeira/assets/fonts').rglob('*.ttf'))
            if sys.platform == 'win32':
                sessao = nuvem.Sessao('ficticio', 'teste@example.invalid', 'acesso-ficticio', 'renovavel-ficticio', 1)
                projeto = 'https://pacote-ficticio.supabase.co'
                sessao_windows.salvar(pasta, projeto, sessao)
                assert sessao_windows.carregar(pasta, projeto).refresh_token == sessao.refresh_token
                sessao_windows.remover(pasta)
            janela._fechamento_autorizado = True
            janela.close(); janela.deleteLater(); app.processEvents()
            janela = None
            resultado.update(sucesso=True, sqlite=sqlite3.sqlite_version, banco_isolado=True,
                backup_seguranca=True, cadastro=True, parcelas=True, excel=True,
                navegacao=True, atualizador_real=True, sem_rede=True)
    except Exception as erro:
        # O relatório público não contém argumentos, credenciais ou mensagens de exceção.
        resultado['erro_tipo'] = type(erro).__name__
        resultado['pilha'] = [dict(arquivo=Path(f.filename).name, funcao=f.name, linha=f.lineno)
            for f in traceback.extract_tb(erro.__traceback__)]
    finally:
        if janela is not None:
            janela._fechamento_autorizado = True
            janela.close()
        Path(relatorio).write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if resultado['sucesso'] else 1
