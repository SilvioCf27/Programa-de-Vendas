"""Equivalência da primeira fatia Qt em bancos temporários, sem conta ou rede."""
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from datetime import date
from decimal import Decimal
import sqlite3
import csv
from contextlib import closing

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'prototipos/reforma'))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from PySide6 import QtCore as Q, QtWidgets as W
from noturno_demo import preparar as preparar_qt
from reforma_desktop.vendas import Consulta, ServicoVendas, valor
from reforma_desktop.aplicacao import Janela, FormularioVenda
from reforma_desktop.ambiente import preparar as preparar_pasta
from ProgramFeira.database import banco
from ProgramFeira import fechamento
from ProgramFeira.exportacao import exportar


class VendasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = preparar_qt()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.pasta = Path(self.temp.name)
        self.patch = patch.object(banco, 'CAMINHO_BANCO', self.pasta / 'vendas.db')
        self.patch.start()
        self.addCleanup(self.patch.stop)
        self.servico = ServicoVendas(lambda: date(2026, 10, 27))
        self.addCleanup(fechamento.definir, {'versao': 0, 'historico': []})

    def criar(self, dia, dinheiro='10', pix='20', cartao='30'):
        self.servico.salvar(dia, dinheiro, pix, cartao)
        return banco.buscar_venda_por_data(dia)

    def janela(self):
        j = Janela(self.servico, self.pasta)
        self.addCleanup(j.deleteLater)
        self.addCleanup(j.close)
        j.show()
        self.app.processEvents()
        return j

    def test_cadastro_duplicado_nao_bloqueia_proxima_gravacao(self):
        self.criar('2026-10-23')
        with self.assertRaisesRegex(ValueError, 'Já existe'):
            self.criar('2026-10-23')
        self.criar('2026-10-25')
        self.assertEqual(len(banco.listar_vendas()), 2)

    def test_valores_data_e_dias_zerados(self):
        self.criar('2026-10-23', '1.234,56', '2,34', '')
        self.criar('2026-10-25', '', '0', '0')
        r = self.servico.consultar(Consulta())
        self.assertEqual(r['indicadores']['quantidade'], 2)
        self.assertEqual(r['indicadores']['total'], Decimal('1236.90'))
        for n in ['-1', 'NaN', 'Infinity', 'abc', '1,234', '1e50', '9007199254740993']:
            with self.subTest(n=n), self.assertRaises(ValueError):
                valor(n)
        with self.assertRaises(ValueError):
            self.criar('2026-02-30')

    def test_edicao_por_id_recusa_estado_alterado_e_data_duplicada(self):
        a = self.criar('2026-10-23')
        self.criar('2026-10-25')
        with self.assertRaises(ValueError):
            self.servico.salvar('2026-10-25', '2', '3', '4', a)
        self.servico.salvar('2026-10-23', '100', '0', '0', a)
        with self.assertRaisesRegex(ValueError, 'mudou'):
            self.servico.salvar('2026-10-23', '999', '0', '0', a)
        self.assertEqual(banco.buscar_venda_por_id(a[0])[2], 100)

    def test_fechamento_mes_ano_e_dia_da_semana(self):
        for dia in ['2026-10-09', '2026-10-10', '2026-11-09', '2026-11-10']:
            self.criar(dia)
        self.assertEqual(Consulta(2026, 10, True).limites(), ('2026-10-10', '2026-11-09'))
        self.assertEqual(len(self.servico.consultar(Consulta(2026, 10, True))['linhas']), 2)
        self.assertEqual(len(self.servico.consultar(Consulta(2026, 10))['linhas']), 2)
        self.assertEqual(len(self.servico.consultar(Consulta(2026))['linhas']), 4)
        self.assertEqual(len(self.servico.consultar(Consulta(2026, dias=(4,)))['linhas']), 1)
        fechamento.definir({'versao': 1, 'historico': [{'vigencia': '2026-11-01', 'dia': 14}]})
        self.assertEqual(Consulta(2026, 9, True).limites(), ('2026-09-10', '2026-10-09'))
        self.assertEqual(Consulta(2026, 11, True).limites(), ('2026-11-15', '2026-12-14'))

    def test_semana_independe_do_filtro_e_exclui_futuro(self):
        self.criar('2026-10-26', '0', '0', '0')
        self.criar('2026-10-27')
        self.criar('2026-10-28', '999', '0', '0')
        s = self.servico.consultar(Consulta(2025))['semana']
        self.assertEqual((s['quantidade'], s['total']), (2, Decimal('60')))
        self.servico.hoje = lambda: date(2026, 11, 2)
        self.assertEqual(self.servico.consultar(Consulta())['semana']['quantidade'], 0)

    def test_lote_atomico_e_backup_anterior_verificado(self):
        a, b = self.criar('2026-10-23'), self.criar('2026-10-25')
        self.servico.excluir([a, b])
        self.assertEqual(banco.listar_vendas(), [])
        ult = max((self.pasta / 'backups').glob('*.db'), key=lambda p: p.stat().st_mtime_ns)
        with closing(sqlite3.connect(ult)) as c:
            self.assertEqual(c.execute('PRAGMA integrity_check').fetchone(), ('ok',))
            self.assertEqual(c.execute('SELECT count(*) FROM vendas').fetchone(), (2,))

    def test_lote_conflitante_nao_remove_nenhuma_venda(self):
        a, b = self.criar('2026-10-23'), self.criar('2026-10-25')
        banco.atualizar_venda(b[0], b[1], 99, 0, 0)
        with self.assertRaises(ValueError):
            self.servico.excluir([a, b])
        self.assertEqual(len(banco.listar_vendas()), 2)

    def test_falha_no_backup_impede_cadastro_edicao_e_exclusao(self):
        a = self.criar('2026-10-23')
        with patch.object(banco, 'backup_arquivo', side_effect=OSError('backup indisponível')):
            for acao in [lambda: self.criar('2026-10-25'), lambda: self.servico.salvar(a[1], '999', '', '', a), lambda: self.servico.excluir([a])]:
                with self.assertRaises(OSError):
                    acao()
        self.assertEqual(banco.listar_vendas(), [a])

    def test_filtros_pendentes_nao_alteram_relatorio_e_gravacao(self):
        self.criar('2026-10-23')
        self.criar('2026-09-23')
        j = self.janela()
        antes = j.resultado['relatorio']
        j.mes.setCurrentIndex(9)
        self.assertTrue(j.aviso.text())
        self.assertEqual(j.consulta.mes, 10)
        self.assertIs(j.resultado['relatorio'], antes)
        self.criar('2026-10-25')
        j.atualizar()
        self.assertEqual(j.consulta.mes, 10)
        self.assertEqual(len(j.resultado['linhas']), 2)
        j.aplicar()
        self.assertEqual(j.consulta.mes, 9)
        self.assertEqual(len(j.resultado['linhas']), 1)

    def test_selecao_de_mes_infere_ano_e_limpar_aplica_todos(self):
        j = self.janela()
        j.limpar()
        j.mes.setCurrentIndex(3)
        self.assertEqual(j.ano.currentText(), '2026')
        self.assertFalse(j.fechamento.isChecked())
        j.aplicar()
        self.assertEqual(j.consulta, Consulta(2026, 3))
        j.ano.setCurrentIndex(0)
        self.assertEqual(j.mes.currentIndex(), 0)
        self.assertFalse(j.fechamento.isChecked())

    def test_cancelar_e_corrigir_duplicidade_no_formulario(self):
        self.criar('2026-10-23')
        callback = []
        form = FormularioVenda(self.servico, lambda: callback.append(True))
        self.addCleanup(form.deleteLater)
        form.data.setDate(Q.QDate(2026, 10, 23))
        form.gravar()
        self.assertIn('Já existe', form.erro.text())
        self.assertTrue(form.salvar.isEnabled())
        form.data.setDate(Q.QDate(2026, 10, 25))
        form.gravar()
        self.assertEqual(callback, [True])
        cancel = FormularioVenda(self.servico, lambda: self.fail('Cancelamento gravou'))
        cancel.reject()
        cancel.deleteLater()
        self.assertEqual(len(banco.listar_vendas()), 2)

    def test_paginacao_busca_e_exportacao_da_consulta_completa(self):
        for dia in range(10, 26):
            self.criar(f'2026-10-{dia:02}')
        j = self.janela()
        j.tabela.todos.setChecked(True)
        self.assertEqual(len(j.tabela.selecionados()), 10)
        j.tabela.mover(1)
        self.assertEqual(j.tabela.selecionados(), [])
        self.assertEqual(len(j.tabela.visiveis), 6)
        j.busca.setText('23/10/2026')
        self.assertEqual(len(j.tabela.linhas), 1)
        for formato in ['csv', 'xlsx']:
            arquivo = self.pasta / ('exportacao.' + formato)
            exportar(arquivo, j.resultado['relatorio'], formato)
            if formato == 'csv':
                with arquivo.open(encoding='utf-8-sig') as f:
                    linhas = list(csv.reader(f, delimiter=';'))
                self.assertEqual(sum(len(r) == 7 and r[0].isdigit() for r in linhas), 16)
            else:
                from openpyxl import load_workbook
                livro = load_workbook(arquivo)
                self.assertEqual(livro['Vendas'].max_row, 17)
                livro.close()

    def test_ordenacao_numerica_preserva_id_da_acao(self):
        a = self.criar('2026-10-23', '2', '0', '0')
        b = self.criar('2026-10-25', '100', '0', '0')
        j = self.janela()
        j.ordenar(4)
        self.assertEqual([r[0] for r in j.tabela.linhas], [a[0], b[0]])
        j.ordenar(4)
        self.assertEqual([r[0] for r in j.tabela.linhas], [b[0], a[0]])

    def test_navegacao_nao_reconstroi_vendas_nem_repete_timer(self):
        j = self.janela()
        tabela, timer = j.tabela, j.relogio.timerId()
        for _ in range(10):
            for nome in ['Análises', 'Finanças', 'Vendas']:
                j.mostrar(nome)
        self.assertIs(j.tabela, tabela)
        self.assertEqual(j.relogio.timerId(), timer)
        self.assertEqual(len(j.findChildren(Q.QTimer)), 1)
        self.assertNotIn('ProgramFeira.main', sys.modules)

    def test_virada_do_dia_atualiza_semana_sem_aplicar_filtro_pendente(self):
        self.criar('2026-10-27')
        j = self.janela()
        j.mes.setCurrentIndex(9)
        j.servico.hoje = lambda: date(2026, 11, 2)
        j.verificar_dia()
        self.assertEqual(j.resultado['semana']['quantidade'], 0)
        self.assertEqual(j.consulta.mes, 10)


class AmbienteTests(unittest.TestCase):
    def test_recusa_pasta_pessoal_e_pasta_com_banco_sem_marcador(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {'PROGRAMFEIRA_DATA_DIR': str(Path(temp) / 'pessoal')}):
            for caminho in [Path(temp), Path(temp) / 'pessoal', Path(temp) / 'pessoal/subpasta']:
                with self.assertRaises(ValueError):
                    preparar_pasta(caminho)
            estudo = Path(temp) / 'estudo'
            estudo.mkdir()
            (estudo / 'vendas.db').write_bytes(b'nao tocar')
            with self.assertRaises(ValueError):
                preparar_pasta(estudo)
            self.assertEqual((estudo / 'vendas.db').read_bytes(), b'nao tocar')

    def test_pasta_vazia_separada_pode_ser_reaberta(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(os.environ, {'PROGRAMFEIRA_DATA_DIR': str(Path(temp) / 'pessoal')}):
            estudo = Path(temp) / 'estudo'
            preparar_pasta(estudo)
            self.assertEqual(os.environ['PROGRAMFEIRA_DATA_DIR'], str(estudo.resolve()))
            # Uma segunda execução inicia com o ambiente normal do sistema.
            os.environ['PROGRAMFEIRA_DATA_DIR'] = str(Path(temp) / 'pessoal')
            self.assertEqual(preparar_pasta(estudo), estudo.resolve())
