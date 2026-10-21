from contextlib import closing
from dados_ficticios import criar_banco
import csv
import hashlib
import sqlite3
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from ProgramFeira.exportacao import exportar, montar_relatorio

ROOT = Path(__file__).resolve().parents[1]


class ExportacaoTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.origem = criar_banco(self.root / 'origem.db')
        with closing(sqlite3.connect(self.origem)) as c:
            self.vendas = c.execute('SELECT id,data,dinheiro,pix,cartao FROM vendas ORDER BY data DESC').fetchall()
        self.relatorio = montar_relatorio(self.vendas, 'Todas as vendas', 'Sem comparação')

    def test_csv_9_registros_centavos_totais_e_acentos(self):
        p = self.root / 'vendas.csv'
        exportar(p, self.relatorio, 'csv')
        self.assertTrue(p.read_bytes().startswith(b'\xef\xbb\xbf'))
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f, delimiter=';'))
        self.assertEqual(rows[0][5], 'Cartão')
        self.assertEqual(len(rows[1:10]), 9)
        total = sum(Decimal(row[6].replace(',', '.')) for row in rows[1:10])
        self.assertEqual(total, Decimal('561.25'))
        self.assertIn(['Total geral', '561,25'], rows)
        self.assertIn(['Registros / dias de feira', '9'], rows)

    def test_excel_datas_numericas_duas_abas_e_comparacao(self):
        p = self.root / 'vendas.xlsx'
        exportar(p, self.relatorio, 'xlsx')
        w = load_workbook(p, data_only=False)
        try:
            self.assertEqual(w.sheetnames, ['Vendas', 'Resumo'])
            self.assertEqual(w['Vendas'].max_row, 10)
            self.assertIsInstance(w['Vendas']['B2'].value, datetime)
            self.assertEqual(w['Vendas']['B2'].number_format, 'dd/mm/yyyy')
            totais = [Decimal(str(row[0])) for row in w['Vendas'].iter_rows(min_row=2, min_col=7, max_col=7, values_only=True)]
            self.assertEqual(sum(totais), Decimal('561.25'))
            resumo = dict(w['Resumo'].iter_rows(min_row=2, values_only=True))
            self.assertEqual(Decimal(str(resumo['Total geral'])), Decimal('561.25'))
            self.assertEqual(resumo['Comparação com período anterior'], 'Sem comparação')
        finally:
            w.close()

    def test_exportacao_filtrada_e_dia_zerado(self):
        vendas = [v for v in self.vendas if v[1].startswith('2027-')]
        relatorio = montar_relatorio(vendas, 'Ano: 2027', 'Comparação de teste')
        self.assertEqual(len(vendas), 1)
        p = self.root / 'ano.xlsx'
        exportar(p, relatorio, 'xlsx')
        w = load_workbook(p)
        try:
            self.assertEqual(w['Vendas'].max_row, 2)
            self.assertEqual(w['Vendas']['G2'].value, 0)
            self.assertEqual(dict(w['Resumo'].iter_rows(min_row=2, values_only=True))['Filtro aplicado'], 'Ano: 2027')
        finally:
            w.close()

    def test_vazio_exporta_cabecalho_e_resumo_zero(self):
        p = self.root / 'vazio.csv'
        exportar(p, montar_relatorio([], 'Ano sem registros', 'Sem registros'), 'csv')
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f, delimiter=';'))
        self.assertEqual(rows[1], [])
        self.assertIn(['Total geral', '0,00'], rows)

    def test_falha_preserva_arquivo_anterior_e_remove_temporario(self):
        p = self.root / 'preservado.csv'
        p.write_bytes(b'original')
        with patch('ProgramFeira.exportacao.escrever_csv', side_effect=OSError('arquivo bloqueado')):
            with self.assertRaises(OSError):
                exportar(p, self.relatorio, 'csv')
        self.assertEqual(p.read_bytes(), b'original')
        self.assertEqual(list(self.root.glob('.programfeira-*')), [])

    def test_nao_permite_exportar_por_cima_de_banco(self):
        p = self.root / 'vendas.db'
        p.write_bytes(b'preservado')
        with self.assertRaises(ValueError):
            exportar(p, self.relatorio, 'csv')
        self.assertEqual(p.read_bytes(), b'preservado')

    def test_texto_nao_vira_formula(self):
        relatorio = montar_relatorio([(1, '=1+1', 0, 0, 0)], '=1+1', '=1+1')
        p = self.root / 'seguro.xlsx'
        exportar(p, relatorio, 'xlsx')
        w = load_workbook(p)
        try:
            self.assertEqual(w['Vendas']['B2'].data_type, 's')
            self.assertEqual(w['Resumo']['B2'].data_type, 's')
        finally:
            w.close()

    def test_nao_altera_banco_origem(self):
        p = self.origem
        before = hashlib.sha256(p.read_bytes()).digest()
        exportar(self.root / 'vendas.csv', self.relatorio, 'csv')
        self.assertEqual(hashlib.sha256(p.read_bytes()).digest(), before)
